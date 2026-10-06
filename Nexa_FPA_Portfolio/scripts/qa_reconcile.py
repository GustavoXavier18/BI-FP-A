"""QA independente dos valores armazenados no Excel. Não executa o engine DAX/M nem altera a fonte."""
from pathlib import Path
import openpyxl,collections,json,csv,hashlib,sys,math
ROOT=Path(__file__).resolve().parents[1]
p=Path(sys.argv[1]) if len(sys.argv)>1 else ROOT/'data'/'Projeto_FPA_Nexa_Distribuicao.xlsx'
w=openpyxl.load_workbook(p,data_only=True);wf=openpyxl.load_workbook(p,data_only=False)
TOL=.01;TESTS=[]
def rows(n):
 s=list(w[n].values);return [dict(zip(s[0],r)) for r in s[1:] if any(x is not None for x in r)]
def test(name,context,error,limit=TOL):
 TESTS.append({'controle':name,'contexto':context,'erro_absoluto':abs(error),'limite':limit,'status':'OK' if abs(error)<limit else 'REVISAR'})
m=rows('12_MODELO');sales=rows('09_VENDAS');exp=rows('10_DESPESAS');journal=rows('13_DIARIO');pv=rows('16_PONTE_PRECO_VOLUME');con=rows('15_CONCILIACAO')
index={(r['Data'],r['Cenario']):r for r in m};groups=collections.defaultdict(list)
for r in m:groups[(r['Cenario'],r['Data'].year)].append(r)
for r in m:
 c=r['Data'].strftime('%Y-%m')+' | '+r['Cenario'];groups[(r['Cenario'],r['Data'].year)].sort(key=lambda x:x['Data'])
 for name,error in [('Ativo=Passivo+PL',r['Ativo']-r['Passivo']-r['PL']),('FCO Direto=Indireto',r['FCO_Direto']-r['FCO_Indireto']),('Caixa Inicial+Fluxos=Final',r['CaixaInicial']+r['FCO_Indireto']+r['FCI']+r['FCF']-r['CaixaFinal']),('Receita Líquida',r['ReceitaBruta']-r['Deducoes']-r['ReceitaLiquida']),('Lucro Bruto',r['ReceitaLiquida']-r['CMV']-r['LucroBruto']),('EBITDA',r['LucroBruto']-r['Opex']-r['EBITDA']),('Lucro Líquido',r['EBITDA']-r['Depreciacao']-r['Juros']-r['IRSimulado']-r['LucroLiquido']),('FCL',r['FCO_Indireto']-r['Capex']-r['FC_Livre']),('NCG completa',r['Clientes']+r['Estoque']-r['Fornecedores']-r['SalariosPagar']-r['TributosVendaPagar']-r['IRPagar']-r['NCG']),('Financiamento',r['Captacao']-r['Amortizacao']-r['Dividendos']-r['FCF'])]:test(name,c,error)
 if r['Cenario']=='Realizado':
  for f in ['ReceitaBruta','CMV']:test('Vendas × '+f,c,r[f]-sum(x[f] for x in sales if x['Data']==r['Data']))
  test('Despesas × Opex',c,r['Opex']-sum(x['Valor'] for x in exp if x['Data']==r['Data']))
 if r['Cenario']=='Forecast 6+6' and r['Data'].month<=6:
  real=index[(r['Data'],'Realizado')]
  for f,v in r.items():
   if isinstance(v,(int,float)):test('Forecast H1 × Realizado',c+' | '+f,v-real[f])
for (sc,yr),rs in groups.items():
 for prev,cur in zip(rs,rs[1:]):test('Continuidade Caixa',sc+' | '+cur['Data'].strftime('%Y-%m'),cur['CaixaInicial']-prev['CaixaFinal'])
 if yr==2026:test('Abertura Cenário 2026',sc,rs[0]['CaixaInicial']-next(x['CaixaFinal'] for x in m if x['Cenario']=='Realizado' and x['Data'].year==2025 and x['Data'].month==12))
for r in con:test('Conciliação Patrimonial',r['Data'].strftime('%Y-%m')+' | '+r['ContaID'],r['Diferenca'])
entries=collections.defaultdict(float)
for r in journal:entries[r['Lancamento']]+=r['Debito']-r['Credito']
for id,b in entries.items():test('Partidas Dobradas',id,b)
lookup={(r['Data'].year,r['Data'].month,r['Produto'],r['Regiao'],r['Canal']):r for r in sales}
for r in pv:
 key=(r['Data'].month,r['Produto'],r['Regiao'],r['Canal']);a=lookup[(2024,*key)];b=lookup[(2025,*key)];c=' | '.join(map(str,key))
 test('Ponte PV Fecha',c,r['Receita2024']+r['EfeitoVolume']+r['EfeitoPreco']-r['Receita2025'])
 test('Volume Independente',c,r['EfeitoVolume']-(b['Quantidade']-a['Quantidade'])*a['PrecoUnitario'])
 test('Preço Independente',c,r['EfeitoPreco']-b['Quantidade']*(b['PrecoUnitario']-a['PrecoUnitario']))
 test('Ponte Receita2024 × Vendas',c,r['Receita2024']-a['ReceitaBruta'])
 test('Ponte Receita2025 × Vendas',c,r['Receita2025']-b['ReceitaBruta'])
# Reference outputs: all modeled rows in DRE/BP/indicators, every displayed month and annual aggregation.
actual=groups[('Realizado',2025)];annual={f:sum(r[f] for r in actual) for f in actual[0] if isinstance(actual[0][f],(int,float))};end=actual[-1]
positions={'CaixaFinal','Clientes','Estoque','AtivoCirculante','ImobilizadoBruto','DepAcumulada','AtivoNaoCirculante','Ativo','Fornecedores','SalariosPagar','TributosVendaPagar','IRPagar','DividaCP','PassivoCirculante','DividaLP','PassivoNaoCirculante','Divida','Passivo','Capital','LucrosAcumulados','PL','NCG','DSO','DIO','DPO','DividaLiquida','LiquidezCorrente'}
for sheet in ['02_DRE','03_BALANCO','05_INDICADORES']:
 s=w[sheet]
 for rr in range(8,s.max_row+1):
  f=s.cell(rr,2).value
  if f=='IR sobre lucro (simulado)':f='IRSimulado'
  if f not in actual[0]:continue
  for month in range(1,13):
   v=s.cell(rr,month+2).value
   if isinstance(v,(int,float)):test('Output '+sheet,f+' | '+str(month),v-actual[month-1][f])
  v=s.cell(rr,15).value
  if isinstance(v,(int,float)):
   expect=end[f] if f in positions else annual[f]
   if f=='MargemEBITDA':expect=annual['EBITDA']/annual['ReceitaLiquida']
   if f=='MargemLiquida':expect=annual['LucroLiquido']/annual['ReceitaLiquida']
   test('Output '+sheet,f+' | anual',v-expect)
# DFC has sign conventions and opening/ending positions.
mapdfc={'LucroLiquido':('LucroLiquido',1),'Depreciacao':('Depreciacao',1),'FCO_Indireto':('FCO_Indireto',1),'Investimentos (capex)':('Capex',-1),'Captacao':('Captacao',1),'Amortizacao':('Amortizacao',-1),'Dividendos':('Dividendos',-1),'Fluxo de financiamento':('FCF',1),'VariacaoCaixa':('VariacaoCaixa',1),'CaixaInicial':('CaixaInicial',1),'CaixaFinal':('CaixaFinal',1),'CheckDFC':('CheckDFC',1)}
for rr in range(8,w['04_DFC'].max_row+1):
 label=w['04_DFC'].cell(rr,2).value
 if label not in mapdfc:continue
 f,sign=mapdfc[label]
 for mm in range(1,13):test('Output DFC',label+' | '+str(mm),w['04_DFC'].cell(rr,mm+2).value-sign*actual[mm-1][f])
 value=w['04_DFC'].cell(rr,15).value;expect=actual[0][f] if f=='CaixaInicial' else end[f] if f=='CaixaFinal' else sum(r[f] for r in actual)
 test('Output DFC',label+' | anual',value-sign*expect)
summary=[]
for (sc,yr),rs in groups.items():
 rs.sort(key=lambda x:x['Data']);e=rs[-1];s={'Cenario':sc,'Ano':yr}
 for f in ['ReceitaLiquida','EBITDA','LucroLiquido','FCO_Indireto','FC_Livre','Capex','Captacao','Amortizacao','Dividendos']:s[f]=sum(r[f] for r in rs)
 for f in ['CaixaFinal','NCG','Divida','DividaLiquida','DSO','DIO','DPO']:s[f]=e[f]
 s['MargemEBITDA']=s['EBITDA']/s['ReceitaLiquida'];s['CaixaMinimo']=min(r['CaixaFinal'] for r in rs);s['MesesCaixaNegativo']=[r['Data'].strftime('%Y-%m') for r in rs if r['CaixaFinal']<0];summary.append(s)
formula_count=0;cached_core_missing=[];excel_errors=[]
for sheet in wf:
 for row in sheet:
  for c in row:
   if c.data_type=='f':
    formula_count+=1
    if sheet.title=='12_MODELO' and w[sheet.title][c.coordinate].value is None:cached_core_missing.append(c.coordinate)
   if w[sheet.title][c.coordinate].data_type=='e':excel_errors.append((sheet.title,c.coordinate))
out={'source':p.name,'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'tolerance':TOL,'scope':'Stored Excel values + independent identities; no Excel full recalculation or DAX execution.','test_count':len(TESTS),'failed_tests':[t for t in TESTS if t['status']!='OK'],'max_absolute_error':max(t['erro_absoluto'] for t in TESTS),'formula_count':formula_count,'core_missing_cached':cached_core_missing,'excel_errors':excel_errors,'summary':summary}
(ROOT/'docs'/'qa_resultados.json').write_text(json.dumps(out,ensure_ascii=False,indent=2))
with (ROOT/'docs'/'qa_detalhado.csv').open('w',encoding='utf-8-sig',newline='') as f:
 cw=csv.DictWriter(f,fieldnames=list(TESTS[0]));cw.writeheader();cw.writerows(TESTS)
print(json.dumps({k:v for k,v in out.items() if k not in ['summary','failed_tests']},indent=2));print('FAILED',out['failed_tests'][:10])
if out['failed_tests'] or cached_core_missing or excel_errors:sys.exit(1)
